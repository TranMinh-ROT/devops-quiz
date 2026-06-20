#!/usr/bin/env python3
"""
Process CCNA Part7–Part9 DOCX files into structured quiz questions.

Pipeline:
  1. Extract paragraphs from CCNA DOCX files
  2. Split into question blocks by 'Question N.N' headers
  3. Filter out IPv6 and Drag-and-Drop questions
  4. Parse question text, options, answers, explanations
  5. Validate schema (id, question, options, correct answers)
  6. Deduplicate against existing question bank + within CCNA set
  7. Output temp/ccna-parsed-questions.json + temp/ccna-import-report.json
"""
import os
import re
import json
import hashlib
from docx import Document

# ============================================================
# Configuration
# ============================================================

CCNA_SOURCE_DIR = "source-documents/ccna"
CCNA_FILES = ["CCNA_Part7.docx", "CCNA_Part8.docx", "CCNA_Part9.docx"]
EXISTING_QUIZ_DIR = "src/data"
OUTPUT = "temp/ccna-parsed-questions.json"
REPORT = "temp/ccna-import-report.json"

# Filtering keywords (all checked case-insensitively unless noted)
IPV6_KEYWORDS = [
    "ipv6", "2001:", "fe80", "fc00", "eui-64",
    "link-local", "unique local address", "neighbor discovery",
    "slaac",
]
# 'ND' needs word-boundary matching to avoid false positives
IPV6_REGEX_PATTERNS = [re.compile(r'\bND\b')]

DRAG_DROP_KEYWORDS = [
    "drag and drop", "drag-and-drop", "drag the",
]

# ============================================================
# Step 1: Extract paragraphs
# ============================================================

def extract_paragraphs(fpath):
    """Extract non-empty paragraphs from a DOCX file."""
    doc = Document(fpath)
    return [(i, p.text.strip()) for i, p in enumerate(doc.paragraphs) if p.text.strip()]


# ============================================================
# Step 2: Split into question blocks
# ============================================================

def split_into_blocks(paragraphs, fname):
    """Split paragraphs into question blocks using 'Question N.N' headers."""
    blocks = []
    current = None
    for idx, text in paragraphs:
        m = re.match(r'^Question\s+(\d+(?:\.\d+)?)\s*$', text)
        if m:
            if current is not None:
                blocks.append(current)
            current = {
                "sourceFile": fname,
                "sourceQuestionNumber": m.group(1),
                "lines": []
            }
            continue
        if current is not None:
            current["lines"].append(text)
    if current is not None:
        blocks.append(current)
    return blocks


# ============================================================
# Step 3: Filtering
# ============================================================

def is_ipv6_question(block):
    """Check if a question block contains IPv6-related content."""
    full_text = "\n".join(block["lines"])
    full_lower = full_text.lower()
    for kw in IPV6_KEYWORDS:
        if kw in full_lower:
            return True
    for pat in IPV6_REGEX_PATTERNS:
        if pat.search(full_text):
            return True
    return False


def is_drag_drop_question(block):
    """Check if a question block is a Drag-and-Drop question."""
    full_text = "\n".join(block["lines"]).lower()
    for kw in DRAG_DROP_KEYWORDS:
        if kw in full_text:
            return True
    return False


def is_removed_placeholder(block):
    """Check if a question was explicitly marked as removed/duplicated."""
    full_text = "\n".join(block["lines"]).lower()
    return "duplicated so we removed" in full_text or "this question is duplicated" in full_text


# ============================================================
# Step 4: Parse question blocks
# ============================================================

def parse_question_block(block):
    """Parse a single question block into structured data."""
    lines = block["lines"]
    source_file = block["sourceFile"]
    source_qnum = block["sourceQuestionNumber"]
    issues = []

    if not lines:
        return None, ["EMPTY_BLOCK"]

    # Find the Answer: line
    answer_idx = -1
    for i, line in enumerate(lines):
        if re.match(r'^Answer:\s*', line):
            answer_idx = i
            break

    if answer_idx == -1:
        return None, ["NO_ANSWER_LINE"]

    # --- Parse question text ---
    # Everything before Answer: line that isn't an option is question text
    # Options start with A. B. C. D. (or are in a single paragraph with newlines)
    question_lines = []
    option_lines = []
    explanation_lines = []

    # First, separate pre-answer content into question text vs options
    pre_answer = lines[:answer_idx]
    for line in pre_answer:
        # Check if this line contains labeled options (A. B. C. D.)
        # Options may be in a single paragraph separated by \n
        sub_lines = line.split("\n")
        has_option = False
        for sl in sub_lines:
            sl_stripped = sl.strip()
            if re.match(r'^[A-F][\.\)]\s*', sl_stripped) and len(sl_stripped) <= 2 or re.match(r'^[A-F][\.\)]\s+', sl_stripped):
                has_option = True
                break

        if has_option:
            # This paragraph contains options - split and add
            for sl in sub_lines:
                sl_stripped = sl.strip()
                if sl_stripped:
                    option_lines.append(sl_stripped)
        else:
            question_lines.append(line)

    # --- Parse answer ---
    answer_line = lines[answer_idx]
    answer_text = answer_line[len("Answer:"):].strip()
    correct_answers = re.findall(r'[A-F]', answer_text)

    # --- Parse post-answer content (explanations, references) ---
    post_answer = lines[answer_idx + 1:]
    explanation_text = ""
    references = []
    in_explanation = False

    for line in post_answer:
        if line.strip() == "Explanation":
            in_explanation = True
            continue
        if line.strip().startswith("Reference:") or line.strip().startswith("https://"):
            urls = re.findall(r'https?://[^\s,\'"]+', line)
            references.extend(urls)
            continue
        if in_explanation:
            if explanation_text:
                explanation_text += " " + line.strip()
            else:
                explanation_text = line.strip()

    # --- Build question text ---
    question_text = " ".join(q.strip() for q in question_lines if q.strip())

    # --- Parse options ---
    options = parse_options(option_lines)

    # --- Detect question type ---
    if len(correct_answers) > 1:
        q_type = "multiple"
        required_count = len(correct_answers)
    else:
        q_type = "single"
        required_count = 1

    # --- Validate ---
    if not question_text:
        issues.append("MISSING_QUESTION_TEXT")
    if not options:
        issues.append("NO_OPTIONS")
    if not correct_answers:
        issues.append("NO_CORRECT_ANSWERS")

    # Validate correct answers reference valid option IDs
    opt_ids = {o["id"] for o in options}
    for ca in correct_answers:
        if ca not in opt_ids:
            issues.append(f"INVALID_ANSWER_ID: {ca}")

    if len(options) < 2:
        issues.append("TOO_FEW_OPTIONS")

    return {
        "sourceFile": source_file,
        "sourceQuestionNumber": source_qnum,
        "type": q_type,
        "requiredAnswerCount": required_count,
        "questionEnglish": question_text,
        "questionVietnamese": "",
        "options": options,
        "correctAnswers": correct_answers,
        "explanationVietnamese": "",
        "explanationEnglish": explanation_text,
        "optionExplanationsVietnamese": {},
        "references": references,
        "reviewStatus": "clean",
        "reviewReasons": [],
        "sourceFormat": "ccna-docx",
        "_issues": issues,
    }, issues


def parse_options(option_lines):
    """Parse option lines into structured options."""
    options = []
    current_id = None
    current_text_parts = []

    for line in option_lines:
        # Check if this starts a new option
        # Handles both "A. text" and "A. \n..." (empty text, content on next lines)
        m = re.match(r'^([A-F])[\.\)]\s*(.*)', line)
        if m:
            # Save previous option if any
            if current_id is not None:
                options.append({
                    "id": current_id,
                    "textEnglish": "\n".join(p for p in current_text_parts if p),
                    "textVietnamese": "",
                })
            current_id = m.group(1)
            current_text_parts = [m.group(2).strip()]
        elif current_id is not None:
            # Continuation of current option (multi-line config blocks etc.)
            current_text_parts.append(line)

    # Save last option
    if current_id is not None:
        options.append({
            "id": current_id,
            "textEnglish": "\n".join(p for p in current_text_parts if p),
            "textVietnamese": "",
        })

    return options


# ============================================================
# Step 5: Duplicate detection
# ============================================================

def normalize_text(text):
    """Normalize text for duplicate detection."""
    if not text:
        return ""
    t = text.lower().strip()
    t = re.sub(r'\s+', ' ', t)
    t = re.sub(r'[^\w\s]', '', t)
    return t


def text_signature(question_text, options):
    """Create a signature for duplicate detection."""
    norm_q = normalize_text(question_text)
    norm_opts = sorted(normalize_text(o.get("textEnglish", "")) for o in options)
    combined = norm_q + "|" + "|".join(norm_opts)
    return hashlib.md5(combined.encode()).hexdigest()


def load_existing_questions():
    """Load existing quiz questions for duplicate checking."""
    existing = []
    for fname in os.listdir(EXISTING_QUIZ_DIR):
        if fname.startswith("quiz-") and fname.endswith(".json"):
            fpath = os.path.join(EXISTING_QUIZ_DIR, fname)
            with open(fpath, 'r', encoding='utf-8') as f:
                data = json.load(f)
            existing.extend(data.get("questions", []))
    return existing


def is_similar_question(q1_text, q2_text, threshold=0.85):
    """Check if two questions are similar enough to be considered duplicates."""
    words1 = set(normalize_text(q1_text).split())
    words2 = set(normalize_text(q2_text).split())
    if not words1 or not words2:
        return False
    intersection = words1 & words2
    union = words1 | words2
    jaccard = len(intersection) / len(union)
    return jaccard >= threshold


# ============================================================
# Main Pipeline
# ============================================================

def main():
    print("=" * 60)
    print("CCNA Import Pipeline")
    print("=" * 60)

    # --- Step 1: Extract ---
    all_blocks = []
    file_stats = {}

    for fname in CCNA_FILES:
        fpath = os.path.join(CCNA_SOURCE_DIR, fname)
        if not os.path.exists(fpath):
            print(f"WARNING: {fpath} not found, skipping")
            continue

        paragraphs = extract_paragraphs(fpath)
        blocks = split_into_blocks(paragraphs, fname)
        file_stats[fname] = {"total_blocks": len(blocks)}
        all_blocks.extend(blocks)
        print(f"  {fname}: {len(paragraphs)} paragraphs, {len(blocks)} question blocks")

    total_discovered = len(all_blocks)
    print(f"\nTotal questions discovered: {total_discovered}")

    # --- Step 2: Filter removed placeholders ---
    removed_placeholder_count = 0
    filtered_blocks = []
    for block in all_blocks:
        if is_removed_placeholder(block):
            removed_placeholder_count += 1
            print(f"  Skipping removed placeholder: {block['sourceFile']} Q{block['sourceQuestionNumber']}")
        else:
            filtered_blocks.append(block)

    # --- Step 3: Filter IPv6 ---
    ipv6_removed = []
    non_ipv6_blocks = []
    for block in filtered_blocks:
        if is_ipv6_question(block):
            ipv6_removed.append(f"{block['sourceFile']} Q{block['sourceQuestionNumber']}")
        else:
            non_ipv6_blocks.append(block)

    print(f"\nIPv6 questions removed: {len(ipv6_removed)}")
    for q in ipv6_removed:
        print(f"  - {q}")

    # --- Step 4: Filter Drag-and-Drop ---
    drag_removed = []
    clean_blocks = []
    for block in non_ipv6_blocks:
        if is_drag_drop_question(block):
            drag_removed.append(f"{block['sourceFile']} Q{block['sourceQuestionNumber']}")
        else:
            clean_blocks.append(block)

    print(f"\nDrag-and-Drop questions removed: {len(drag_removed)}")
    for q in drag_removed:
        print(f"  - {q}")

    # --- Step 5: Parse ---
    parsed_questions = []
    malformed = []
    for block in clean_blocks:
        q, issues = parse_question_block(block)
        if q is None or any(i in ["NO_OPTIONS", "NO_CORRECT_ANSWERS", "NO_ANSWER_LINE", "EMPTY_BLOCK", "TOO_FEW_OPTIONS"] for i in issues):
            malformed.append({
                "source": block["sourceFile"],
                "qnum": block["sourceQuestionNumber"],
                "issues": issues,
            })
        else:
            parsed_questions.append(q)

    print(f"\nMalformed questions removed: {len(malformed)}")
    for m in malformed:
        print(f"  - {m['source']} Q{m['qnum']}: {m['issues']}")

    # --- Step 6: Deduplicate ---
    # Load existing questions
    existing_questions = load_existing_questions()
    existing_sigs = set()
    for eq in existing_questions:
        sig = text_signature(eq.get("questionEnglish", ""), eq.get("options", []))
        existing_sigs.add(sig)

    print(f"\nExisting question bank: {len(existing_questions)} questions ({len(existing_sigs)} unique signatures)")

    # Deduplicate against existing
    dupes_existing = []
    non_dupe_questions = []
    for q in parsed_questions:
        sig = text_signature(q["questionEnglish"], q["options"])
        if sig in existing_sigs:
            dupes_existing.append(f"{q['sourceFile']} Q{q['sourceQuestionNumber']}")
        else:
            non_dupe_questions.append(q)

    print(f"Duplicates against existing bank: {len(dupes_existing)}")
    for d in dupes_existing:
        print(f"  - {d}")

    # Deduplicate within CCNA set
    dupes_internal = []
    seen_sigs = set()
    unique_questions = []
    for q in non_dupe_questions:
        sig = text_signature(q["questionEnglish"], q["options"])
        if sig in seen_sigs:
            dupes_internal.append(f"{q['sourceFile']} Q{q['sourceQuestionNumber']}")
        else:
            seen_sigs.add(sig)
            unique_questions.append(q)

    print(f"Internal duplicates removed: {len(dupes_internal)}")
    for d in dupes_internal:
        print(f"  - {d}")

    # Also check for fuzzy duplicates within CCNA set
    fuzzy_dupes = []
    final_questions = []
    for i, q in enumerate(unique_questions):
        is_fuzzy_dupe = False
        for j, existing_q in enumerate(final_questions):
            if is_similar_question(q["questionEnglish"], existing_q["questionEnglish"]):
                fuzzy_dupes.append(f"{q['sourceFile']} Q{q['sourceQuestionNumber']} (similar to Q{existing_q['sourceQuestionNumber']})")
                is_fuzzy_dupe = True
                break
        if not is_fuzzy_dupe:
            # Also check against existing bank
            is_fuzzy_existing = False
            for eq in existing_questions:
                if is_similar_question(q["questionEnglish"], eq.get("questionEnglish", "")):
                    fuzzy_dupes.append(f"{q['sourceFile']} Q{q['sourceQuestionNumber']} (similar to existing ID {eq['id']})")
                    is_fuzzy_existing = True
                    break
            if not is_fuzzy_existing:
                final_questions.append(q)

    print(f"Fuzzy duplicates removed: {len(fuzzy_dupes)}")
    for d in fuzzy_dupes:
        print(f"  - {d}")

    # --- Step 7: Assign IDs and build option explanations ---
    # Find max existing ID
    max_existing_id = max((q["id"] for q in existing_questions), default=0)
    print(f"\nMax existing question ID: {max_existing_id}")

    for i, q in enumerate(final_questions):
        q["id"] = max_existing_id + i + 1

        # Build optionExplanationsVietnamese structure for compatibility
        opt_expls = {}
        for opt in q["options"]:
            opt_expls[opt["id"]] = {
                "isCorrect": opt["id"] in q["correctAnswers"],
                "explanation": ""
            }
        q["optionExplanationsVietnamese"] = opt_expls

        # Remove internal _issues field
        del q["_issues"]

    # --- Summary ---
    final_count = len(final_questions)
    single_count = sum(1 for q in final_questions if q["type"] == "single")
    multi_count = sum(1 for q in final_questions if q["type"] == "multiple")

    print(f"\n{'=' * 60}")
    print(f"IMPORT SUMMARY")
    print(f"{'=' * 60}")
    print(f"Total questions discovered:        {total_discovered}")
    print(f"Removed placeholders:              {removed_placeholder_count}")
    print(f"IPv6 questions removed:            {len(ipv6_removed)}")
    print(f"Drag-and-Drop questions removed:   {len(drag_removed)}")
    print(f"Malformed questions removed:        {len(malformed)}")
    print(f"Duplicates (existing bank):         {len(dupes_existing)}")
    print(f"Duplicates (internal exact):        {len(dupes_internal)}")
    print(f"Duplicates (fuzzy):                 {len(fuzzy_dupes)}")
    print(f"Final imported count:              {final_count}")
    print(f"  Single-answer:                   {single_count}")
    print(f"  Multi-answer:                    {multi_count}")
    print(f"  ID range:                        {final_questions[0]['id']} – {final_questions[-1]['id']}" if final_questions else "  (none)")

    # --- Save output ---
    os.makedirs("temp", exist_ok=True)

    with open(OUTPUT, 'w', encoding='utf-8') as f:
        json.dump(final_questions, f, indent=2, ensure_ascii=False)
    print(f"\nSaved {final_count} questions to {OUTPUT}")

    # Save report
    report = {
        "totalDiscovered": total_discovered,
        "removedPlaceholders": removed_placeholder_count,
        "ipv6Removed": len(ipv6_removed),
        "ipv6RemovedList": ipv6_removed,
        "dragDropRemoved": len(drag_removed),
        "dragDropRemovedList": drag_removed,
        "malformedRemoved": len(malformed),
        "malformedDetails": malformed,
        "duplicatesExisting": len(dupes_existing),
        "duplicatesExistingList": dupes_existing,
        "duplicatesInternal": len(dupes_internal),
        "duplicatesInternalList": dupes_internal,
        "duplicatesFuzzy": len(fuzzy_dupes),
        "duplicatesFuzzyList": fuzzy_dupes,
        "finalImportedCount": final_count,
        "singleAnswerCount": single_count,
        "multiAnswerCount": multi_count,
        "idRange": [final_questions[0]["id"], final_questions[-1]["id"]] if final_questions else [],
        "fileStats": file_stats,
    }

    with open(REPORT, 'w', encoding='utf-8') as f:
        json.dump(report, f, indent=2, ensure_ascii=False)
    print(f"Saved import report to {REPORT}")

    # --- Show sample questions ---
    print(f"\n{'=' * 60}")
    print(f"SAMPLE IMPORTED QUESTIONS")
    print(f"{'=' * 60}")
    for q in final_questions[:3]:
        print(f"\n--- ID {q['id']} (Source: {q['sourceFile']} Q{q['sourceQuestionNumber']}) ---")
        print(f"  Type: {q['type']} (required: {q['requiredAnswerCount']})")
        print(f"  Q: {q['questionEnglish'][:150]}")
        for opt in q["options"]:
            marker = "✓" if opt["id"] in q["correctAnswers"] else " "
            print(f"  [{marker}] {opt['id']}. {opt['textEnglish'][:80]}")
        if q["explanationEnglish"]:
            print(f"  Explanation: {q['explanationEnglish'][:120]}...")

    print(f"\nDone.")


if __name__ == "__main__":
    main()
